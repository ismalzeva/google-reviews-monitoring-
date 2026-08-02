"""Gate L — Live Public Review Pilot & Data Validation (RUN_11).

Automated quality gate for the live pilot path. All vendor calls use HTTP
stubs — NEVER hits the production Outscraper API. Live verification itself
is a separate manual checklist (docs/RUN_11_LIVE_ACTIVATION_CHECKLIST.md).

Covers: live configuration validation, secret redaction, pilot outlet
restriction, Harjamukti exclusion, max review limit, pagination statistics,
incremental second sync, no duplicates after repeat sync, cross-branch
isolation, unknown geography, live source label, dashboard with Outscraper
source, export with live-source metadata, direct reply disabled, auto-reply OFF.
"""
import io
import json
import os
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mktemp(suffix='gate_l.db')}"
os.environ["SECRET_KEY"] = "test-secret"
os.environ["FLASK_ENV"] = "testing"
os.environ["GRM_PILOT_MODE"] = "true"
os.environ["GRM_PILOT_MAX_OUTLETS"] = "1"
os.environ["PUBLIC_REVIEW_PROVIDER"] = "mock"
os.environ.pop("OUTSCRAPER_API_KEY", None)

from werkzeug.security import generate_password_hash

from app import create_app, db
from app.models.entities import (
    User, Business, Outlet, Review, ReviewRawPayload, SyncReport,
)
from app.services.public_provider import build_public_review_adapter
from app.services.feature_flags import (
    is_auto_reply_enabled, pilot_max_outlets, is_outlet_pilot_active,
)
from app.adapters.outscraper_public_review_adapter import (
    OutscraperPublicReviewAdapter, OutscraperError,
)

PLACE_DEPOK = "ChIJ0-depok-margonda-001"
PLACE_MARGONDA = "ChIJ0-depok-margonda-002"
T1_EMAIL = "owner@buburfay.id"
T1_PW = "pass1234"

# ─── FIXTURES (Outscraper-shaped) ──────────────────────────
FIXTURE_LOCATION = {
    "place_id": PLACE_DEPOK,
    "name": "Bubur Fay Depok",
    "full_address": "Jl. Raya Depok No. 21, Kota Depok, Jawa Barat",
    "city": "Depok",
    "state": "Jawa Barat",
    "district": None,
    "rating": 4.5,
    "reviews": 312,
    "latitude": -6.4025,
    "longitude": 106.8187,
    "maps_url": "https://www.google.com/maps/place/?q=place_id:ChIJ0-depok-margonda-001",
}

FIXTURE_REVIEWS = [
    {
        "review_id": "lrv-001",
        "author_title": "Dewi Lestari",
        "review_rating": 5,
        "review_text": "Buburnya enak banget, kuah gurih.",
        "review_datetime_utc": "2026-08-01T07:30:00",
        "owner_answer": "Terima kasih, Kak!",
        "owner_answer_datetime_utc": "2026-08-01T10:00:00",
    },
    {
        "review_id": "lrv-002",
        "author_title": "Budi Hartono",
        "review_rating": 2,
        "review_text": "Porsi kecil.",
        "review_datetime_utc": "2026-07-25T08:10:00",
        "owner_answer": None,
        "owner_answer_datetime_utc": None,
    },
    {
        "review_id": "lrv-003",
        "author_title": "Anonim",
        "review_rating": 1,
        "review_text": "",
        "review_datetime_utc": "2026-07-20T09:00:00",
    },
    {
        # No vendor ID → deterministic hash
        "author_title": "Rina T",
        "review_rating": 4,
        "review_text": "Kebersihan terjaga.",
        "review_datetime_utc": "2026-07-18T06:55:00",
    },
]


class FakeResp:
    def __init__(self, status_code=200, json_data=None, text=None):
        self.status_code = status_code
        self._json = json_data
        self._text = text

    def json(self):
        if self._json is None:
            raise ValueError("No JSON body")
        return self._json

    @property
    def text(self):
        if self._text is not None:
            return self._text
        return json.dumps(self._json) if self._json is not None else ""


def make_fake_get(location=None, reviews=None):
    location = location if location is not None else FIXTURE_LOCATION
    reviews = reviews if reviews is not None else FIXTURE_REVIEWS
    calls = {"n": 0}

    def fake_get(url, params=None, headers=None, timeout=None):
        calls["n"] += 1
        q = (params or {}).get("query")
        rtype = (params or {}).get("reviewsType")
        if rtype == "only_location":
            return FakeResp(json_data={"data": [location]})
        limit = int((params or {}).get("limit", 100))
        offset = int((params or {}).get("offset", 0))
        page = reviews[offset:offset + limit]
        return FakeResp(json_data={"data": [{"place_id": q, "reviews_data": page}]})

    fake_get.calls = calls
    return fake_get


_SEEDED = False


class GateLBase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        global _SEEDED
        cls.app = create_app()
        cls.app.config["TESTING"] = True
        if _SEEDED:
            return
        with cls.app.app_context():
            db.create_all()
            biz1 = Business(id="biz-l-1", tenant_id="tenant-l-1", name="Bubur Fay", brand_name="Bubur Fay")
            u1 = User(id="u-l-1", email=T1_EMAIL, password_hash=generate_password_hash(T1_PW),
                      display_name="Owner", business_id="biz-l-1", role="owner", is_active=True)
            db.session.add_all([biz1, u1])
            db.session.commit()
            # Harjamukti excluded pilot outlet
            db.session.add(Outlet(
                id="outlet-l-hj", tenant_id="tenant-l-1", business_id="biz-l-1",
                name="Bubur Fay Harjamukti", public_place_id="ChIJ0-harjamukti-old-006",
                status="old_or_closed", monitor_enabled=True, reply_enabled=False,
            ))
            db.session.commit()
            _SEEDED = True

    def setUp(self):
        self.client = self.app.test_client()
        self.client.post("/auth/login", data={"email": T1_EMAIL, "password": T1_PW})

    def _make_outlet(self, name="Bubur Fay Depok", place_id=PLACE_DEPOK, status="active"):
        import uuid
        with self.app.app_context():
            o = Outlet(
                id=f"outlet-l-{uuid.uuid4().hex[:10]}",
                tenant_id="tenant-l-1",
                business_id="biz-l-1",
                name=name,
                public_place_id=place_id,
                status=status,
                monitor_enabled=True,
                reply_enabled=False,
            )
            db.session.add(o)
            db.session.commit()
            return o.id

    def _sync_outscraper(self, fake_get, outlet_ids=None, full_sync=False):
        adapter = OutscraperPublicReviewAdapter(api_key="test-key", max_retries=0, page_size=100)
        with mock.patch("app.adapters.outscraper_public_review_adapter.requests.get", fake_get):
            from app.services.sync_service import sync_public_reviews
            with self.app.app_context():
                return sync_public_reviews(
                    tenant_id="tenant-l-1",
                    business_id="biz-l-1",
                    outlet_ids=outlet_ids,
                    adapter=adapter,
                    full_sync=full_sync,
                )


# ─── LIVE CONFIGURATION & SECRETS ──────────────────────────
class TestLiveConfig(GateLBase):
    def test_live_config_requires_key(self):
        os.environ.pop("OUTSCRAPER_API_KEY", None)
        with self.assertRaises(ValueError):
            build_public_review_adapter("outscraper")

    def test_provider_validation(self):
        with self.assertRaises(ValueError):
            build_public_review_adapter("selenium")  # not implemented/unknown

    def test_secret_redaction_in_errors(self):
        # Vendor error must never leak the API key
        secret = "sk-SECRET123"
        a = OutscraperPublicReviewAdapter(api_key=secret, max_retries=0)
        with mock.patch("app.adapters.outscraper_public_review_adapter.requests.get",
                        return_value=FakeResp(status_code=200, text="not json")):
            with self.assertRaises(OutscraperError) as ctx:
                a.list_reviews_by_place_id(PLACE_DEPOK)
        self.assertNotIn(secret, str(ctx.exception))
        self.assertNotIn(secret, ctx.exception.to_dict().get("error", ""))

    def test_missing_key_error_has_no_secret(self):
        os.environ.pop("OUTSCRAPER_API_KEY", None)
        try:
            build_public_review_adapter("outscraper")
        except ValueError as exc:
            self.assertNotIn("sk-", str(exc))


# ─── PILOT RESTRICTIONS ────────────────────────────────────
class TestPilotRestrictions(GateLBase):
    def test_pilot_outlet_restriction(self):
        self.assertEqual(pilot_max_outlets(), 1)
        self.assertLessEqual(pilot_max_outlets(), 1)

    def test_pilot_second_outlet_rejected(self):
        """With GRM_PILOT_MAX_OUTLETS=1, a second active outlet is rejected."""
        with self.app.app_context():
            biz = Business(id="biz-po", tenant_id="tenant-po", name="Pilot Only", brand_name="Pilot Only")
            db.session.add(biz)
            db.session.commit()
            o1 = Outlet(id="po-1", tenant_id="tenant-po", business_id="biz-po",
                        name="Outlet Satu", public_place_id="P1", status="active",
                        monitor_enabled=True, reply_enabled=False)
            db.session.add(o1)
            db.session.commit()
            self.assertTrue(is_outlet_pilot_active(o1))  # first active outlet OK
            # Second active outlet arrives → rejected (limit 1)
            o2 = Outlet(id="po-2", tenant_id="tenant-po", business_id="biz-po",
                        name="Outlet Dua", public_place_id="P2", status="active",
                        monitor_enabled=True, reply_enabled=False)
            db.session.add(o2)
            db.session.commit()
            self.assertFalse(is_outlet_pilot_active(o2))  # limit 1 reached
            self.assertTrue(is_outlet_pilot_active(o1))  # first still active

    def test_harjamukti_exclusion(self):
        with self.app.app_context():
            hj = Outlet.query.filter_by(id="outlet-l-hj").first()
            self.assertFalse(is_outlet_pilot_active(hj))
        rep = self._sync_outscraper(make_fake_get(), outlet_ids=["outlet-l-hj"])
        with self.app.app_context():
            cnt = Review.query.filter_by(outlet_id="outlet-l-hj").count()
        self.assertEqual(cnt, 0)
        # Requested (counted) but skipped — no reviews, no failure
        self.assertEqual(rep["locations_requested"], 1)
        self.assertEqual(rep["locations_succeeded"], 0)
        self.assertEqual(rep["locations_failed"], 0)
        self.assertEqual(rep["reviews_created"], 0)

    def test_max_review_limit(self):
        many = FIXTURE_REVIEWS * 60  # 240
        a = OutscraperPublicReviewAdapter(api_key="k", max_retries=0, page_size=100, max_reviews=2)
        with mock.patch("app.adapters.outscraper_public_review_adapter.requests.get",
                        make_fake_get(reviews=many)):
            items = a.list_reviews_by_place_id(PLACE_DEPOK, limit=100)
        self.assertEqual(len(items), 2)

    def test_pagination_statistics(self):
        # 240 reviews with UNIQUE vendor ids (repeat sync must not conflate them)
        many = [
            dict(FIXTURE_REVIEWS[i % len(FIXTURE_REVIEWS)], review_id=f"lrv-u-{i:04d}")
            for i in range(240)
        ]
        oid = self._make_outlet()
        fake = make_fake_get(reviews=many)
        rep = self._sync_outscraper(fake, outlet_ids=[oid])
        self.assertEqual(rep["reviews_received"], 240)  # multi-page collected
        self.assertEqual(rep["reviews_inserted"], 240)
        self.assertGreater(fake.calls["n"], 2)  # more than one page requested


# ─── SYNC & DATA QUALITY ──────────────────────────────────
class TestLiveDataQuality(GateLBase):
    def test_incremental_second_sync(self):
        oid = self._make_outlet()
        rep1 = self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        rep2 = self._sync_outscraper(make_fake_get(), outlet_ids=[oid])  # incremental
        self.assertEqual(rep1["reviews_inserted"], 4)
        self.assertEqual(rep2["reviews_inserted"], 0)  # no new reviews
        self.assertGreaterEqual(rep2["reviews_skipped"], 1)

    def test_no_duplicate_after_repeat_sync(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid], full_sync=True)
        with self.app.app_context():
            self.assertEqual(Review.query.filter_by(outlet_id=oid).count(), 4)

    def test_cross_branch_isolation(self):
        oid_depok = self._make_outlet(name="Bubur Fay Depok", place_id=PLACE_DEPOK)
        oid_mg = self._make_outlet(name="Bubur Fay Margonda", place_id=PLACE_MARGONDA)
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid_depok])
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid_mg])
        with self.app.app_context():
            depok_ids = {r.source_review_name for r in Review.query.filter_by(outlet_id=oid_depok).all()}
            mg_ids = {r.source_review_name for r in Review.query.filter_by(outlet_id=oid_mg).all()}
        self.assertEqual(len(depok_ids), 4)
        self.assertEqual(len(mg_ids), 4)
        self.assertFalse(depok_ids & mg_ids)  # no cross-branch collision

    def test_unknown_geography_handling(self):
        loc = dict(FIXTURE_LOCATION)
        loc["city"] = None
        loc["state"] = None
        loc["district"] = None
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(location=loc), outlet_ids=[oid])
        with self.app.app_context():
            o = Outlet.query.get(oid)
            self.assertTrue(o.province in (None, "unknown"))
            self.assertTrue(o.city_regency in (None, "unknown"))
            self.assertTrue(o.needs_geographic_resolution)

    def test_live_source_label(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        with self.app.app_context():
            srcs = {r.source for r in Review.query.filter_by(outlet_id=oid).all()}
            rp = ReviewRawPayload.query.filter_by(tenant_id="tenant-l-1").first()
        self.assertEqual(srcs, {"outscraper"})
        self.assertIsNotNone(rp)
        self.assertEqual(rp.source, "outscraper")

    def test_reviewer_masked_after_sync(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        with self.app.app_context():
            names = [r.reviewer_display_name for r in Review.query.filter_by(outlet_id=oid).all()]
        self.assertTrue(all(("***" in n or n == "Anonim") for n in names))
        self.assertIn("Dewi L***", names)

    def test_owner_reply_stored(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        with self.app.app_context():
            rv = Review.query.filter_by(tenant_id="tenant-l-1", outlet_id=oid).filter(
                Review.comment.like("%enak banget%")).first()
        self.assertEqual(rv.owner_reply_text, "Terima kasih, Kak!")

    def test_rating_only_stored(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        with self.app.app_context():
            ro = Review.query.filter_by(tenant_id="tenant-l-1", outlet_id=oid, has_text=False).all()
        self.assertEqual(len(ro), 1)
        self.assertEqual(ro[0].star_rating, 1)

    def test_sync_stats_recorded(self):
        oid = self._make_outlet()
        rep = self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        self.assertEqual(rep["provider"], "outscraper")
        self.assertEqual(rep["sync_status"], "ok")
        with self.app.app_context():
            s = SyncReport.query.filter_by(tenant_id="tenant-l-1").order_by(
                SyncReport.created_at.desc()).first()
            self.assertEqual(s.source, "outscraper")


# ─── DASHBOARD & EXPORT (live source) ──────────────────────
class TestLiveDashboard(GateLBase):
    def _seed_live_data(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        return oid

    def test_dashboard_with_outscraper_source(self):
        self._seed_live_data()
        r = self.client.get("/api/public/analytics/summary?days=90")
        s = r.get_json()
        self.assertEqual(s["total_reviews"], 4)
        self.assertGreaterEqual(s["avg_rating"], 1)
        # category breakdown works on live-source data
        r2 = self.client.get("/api/public/analytics/categories?days=90")
        self.assertEqual(r2.status_code, 200)

    def test_explorer_live_source(self):
        oid = self._seed_live_data()
        r = self.client.get(f"/api/public/analytics/explorer?outlet_id={oid}")
        items = r.get_json()["items"]
        self.assertEqual(len(items), 4)
        self.assertTrue(all(i["source"] == "outscraper" for i in items))
        self.assertTrue(all(("***" in i["reviewer"] or i["reviewer"] == "Anonim") for i in items))

    def test_export_with_live_source_metadata(self):
        import csv as _csv
        self._seed_live_data()
        r = self.client.get("/api/public/analytics/export.csv?city_regency=Depok")
        raw = r.data.decode().splitlines()
        self.assertTrue(any(l.startswith("# kota_kabupaten: Depok") for l in raw))
        parsed = list(_csv.reader(io.StringIO(r.data.decode())))
        # first data row is the header; metadata rows start with '#'
        header_idx = next(i for i, row in enumerate(parsed) if row and not row[0].startswith("#"))
        header = parsed[header_idx]
        source_idx = header.index("sumber")
        for row in parsed[header_idx + 1:]:
            if row:
                self.assertEqual(row[source_idx], "outscraper")

    def test_export_xlsx_live_source(self):
        from openpyxl import load_workbook
        self._seed_live_data()
        r = self.client.get("/api/public/analytics/export.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.data))
        self.assertIn("Export Info", wb.sheetnames)
        self.assertIn("Public Reviews", wb.sheetnames)


# ─── SAFETY ────────────────────────────────────────────────
class TestSafety(GateLBase):
    def test_direct_reply_disabled(self):
        oid = self._make_outlet()
        self._sync_outscraper(make_fake_get(), outlet_ids=[oid])
        with self.app.app_context():
            o = Outlet.query.get(oid)
            self.assertFalse(o.reply_enabled)

    def test_auto_reply_off(self):
        self.assertFalse(is_auto_reply_enabled())

    def test_restart_guide_no_broad_pkill(self):
        from pathlib import Path
        doc = (Path(__file__).resolve().parents[1] / "docs" / "OUTSCRAPER_KEY_OWNER_ACTION.md").read_text()
        # No executable pkill command line (mentions of "pkill" as a *warning*
        # are fine, but the guide must not teach/run it)
        pkill_cmds = [l for l in doc.splitlines() if l.strip().startswith("pkill")]
        self.assertEqual(pkill_cmds, [])
        self.assertIn("run/grm.pid", doc)
        self.assertIn("kill -TERM", doc)
        self.assertIn("venv/bin/python run.py", doc)

    def test_provider_remains_mock_without_key(self):
        """Provider stays mock while OUTSCRAPER_API_KEY is unavailable."""
        os.environ.pop("OUTSCRAPER_API_KEY", None)
        os.environ["PUBLIC_REVIEW_PROVIDER"] = "mock"
        try:
            a = build_public_review_adapter()
            self.assertEqual(type(a).__name__, "MockPublicReviewAdapter")
        finally:
            os.environ.pop("OUTSCRAPER_API_KEY", None)


if __name__ == "__main__":
    unittest.main(verbosity=2)
